from pathlib import Path
from docx import Document
from docx.shared import Inches
root=Path(__file__).resolve().parents[1]
d=Document()
d.add_heading('Sweet Bakery',0)
d.add_paragraph('Mobile Bakery Game — Submission Draft')
d.add_paragraph('Student name: [ADD YOUR NAME]')
d.add_heading('Required submission evidence',1)
d.add_paragraph('GitHub repository: [ADD PUBLIC REPOSITORY URL]')
d.add_paragraph('Successful GitHub Actions run: [ADD RUN URL AND INSERT SCREENSHOT OF ITS GREEN CHECK]')
d.add_paragraph('YouTube gameplay recording: [ADD UNLISTED VIDEO URL AFTER RECORDING THE ANDROID EMULATOR]')
d.add_paragraph('This draft is not ready for LMS submission until the fields above and genuine evidence are supplied. If a video is unavailable, insert actual emulator screenshots; the assignment assigns reduced credit to images alone.')
d.add_heading('Game concept and controls',1)
d.add_paragraph('Sweet Bakery is an English-language offline Android game set in a small bakery. During a 90-second shift, the player selects a customer, adds the exact ordered products to a tray and serves the order before the patience bar runs out. Three lost guests end the shift. Pause and replay are supported.')
d.add_heading('Diversity of meals',1)
d.add_paragraph('The menu includes cupcakes, donuts, croissants, bread, cake, cookies, pie and pretzels. Orders contain different combinations and quantities. Difficulty increases as successful service produces larger orders and faster arrivals.')
d.add_heading('Points management',1)
t=d.add_table(rows=1, cols=2);t.style='Light Shading Accent 1'
t.rows[0].cells[0].text='Event';t.rows[0].cells[1].text='Points / consequence'
for a,b in [('Correct order','+100'),('Fast service (at least half patience remains)','+50 bonus'),('Correct-order streak','+10 per additional order, up to +50'),('Wrong order','−30; streak resets'),('Guest runs out of patience','−50 and one heart; streak resets')]:
 c=t.add_row().cells;c[0].text=a;c[1].text=b
d.add_paragraph('Score cannot fall below zero. A local best score rewards replay. A visible 1,000-point goal awards the Star Baker title without changing the scoring rules. Each service displays feedback explaining the outcome.')
d.add_heading('Visual details and language',1)
d.add_paragraph('The interface uses a rose and cream bakery palette, striped awning, pastry shelves, a warm oven, different customer appearances, illustrated food symbols, patience bars and a serving tray. All game text and instructions are in English.')
d.add_heading('Development previews — browser only',1)
d.add_paragraph('The images below were captured in a mobile-sized desktop browser. They demonstrate the design and must not be presented as Android emulator evidence.')
for filename,label in [('01-welcome.png','Instructions'),('02-gameplay.png','Gameplay and scoring'),('03-shift-complete.png','Shift result')]:
 p=root/'evidence'/'sweet-bakery'/filename
 if p.exists():
  d.add_heading(label,2);d.add_picture(str(p),width=Inches(2.7))
d.save(root/'submission/Sweet-Bakery-Submission-DRAFT.docx')
